import argparse
import collections
import datetime as dt
import hashlib
import gzip
import json
import pathlib
import re

import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('workbook')
args = parser.parse_args()
source = pathlib.Path(args.workbook)
book = openpyxl.load_workbook(source, data_only=True)
formulas = openpyxl.load_workbook(source, data_only=False)
root = pathlib.Path(__file__).resolve().parents[1]
companies = ['景顺长城', '申万菱信', '中信保诚', '中信建投', '国寿安保', '西部利得', '创金合信', '浦银安盛', '国投瑞银', '华泰保兴', '华泰柏瑞', '前海开源', '财通资管', '渤海汇金', '国海富兰克林', '交银施罗德', '兴证全球', '国联安', '易方达', '汇添富', '海富通', '贝莱德', '鑫元', '国泰海通', '广发', '建信', '博时', '兴业', '平安', '万家', '长信', '富国', '大成', '兴银', '鹏扬', '永赢', '中金', '摩根', '中银', '中欧', '华夏', '上银', '国联', '天弘', '银河', '鹏华', '华安', '银华', '华宝', '招商', '泰康', '安信', '嘉实', '长盛', '南方', '信澳', '工银瑞信', '东方红']
aliases = {'兴全': '兴证全球', '浦银': '浦银安盛', '交银': '交银施罗德', '国泰': '国泰海通', '工银': '工银瑞信', '国富': '国海富兰克林'}

def company(name):
    for prefix in sorted(companies + list(aliases), key=len, reverse=True):
        if name.startswith(prefix):
            return aliases.get(prefix, prefix)
    return None

def code_of(value):
    value = str(value or '').strip()
    return value if re.fullmatch(r'\d{6}', value) else None

relations = []
field_definitions = {
    '成立日期': ('established', '成立日期', 'date', ''),
    '基金经理': ('manager', '基金经理', 'text', ''),
    '持有期': ('holding', '原表持有期', 'number', '天'),
    '日回报': ('daily', '日回报', 'number', '%'),
    '近1周': ('week', '近1周', 'number', '%'),
    '近1月': ('month', '近1月', 'number', '%'),
    '近3月': ('quarter', '近3月', 'number', '%'),
    '近6月': ('halfYear', '近6月', 'number', '%'),
    '近1年': ('year', '近1年', 'number', '%'),
    '近3年': ('threeYears', '近3年', 'number', '%'),
    '今年以来': ('ytd', '今年以来', 'number', '%'),
    'YTD收益': ('ytd', '今年以来', 'number', '%'),
    'YTD回撤': ('ytdDrawdown', '年内回撤', 'number', '%'),
    '展示收益（成立以来年化）': ('annualized', '成立以来年化', 'number', '%'),
    '基准': ('benchmark', '原表基准', 'number', '%'),
    '3年基准': ('benchmark3y', '3年基准', 'number', '%'),
    '动态回撤': ('drawdown', '动态回撤', 'number', '%'),
    '近1年最大回撤': ('maxDrawdown1y', '近1年最大回撤', 'number', '%'),
    '近2年最大回撤': ('maxDrawdown2y', '近2年最大回撤', 'number', '%'),
    '基金规模(亿元)': ('size', '基金规模', 'number', '亿元'),
    '基金类型（二级分类）': ('fundType', '基金类型', 'text', ''),
    '管理费': ('managementFee', '管理费率', 'number', '%/年'),
    '托管行': ('custodian', '托管行', 'text', ''),
    '股票仓位': ('stockWeight', '股票仓位', 'number', '%'),
    '转债仓位': ('convertibleWeight', '转债仓位', 'number', '%'),
    '增强策略': ('strategy', '增强策略', 'text', ''),
    '策略': ('riskLabel', '原表策略标签', 'text', ''),
    '标签': ('tag', '原表标签', 'text', ''),
}
field_catalog = {}

def extract_fields(sheet, row):
    header = None
    for index in range(row - 1, 0, -1):
        if sheet.cell(index, 2).value == '基金代码':
            header = index
            break
    result = {}
    if header is None:
        return result
    for column in range(3, sheet.max_column + 1):
        label = sheet.cell(header, column).value
        definition = field_definitions.get(label) if isinstance(label, str) else None
        if isinstance(label, dt.datetime) and column <= 26:
            report_date = label.date().isoformat()
            definition = ('exposure_' + report_date, report_date + ' 折算权益仓位', 'number', '%')
        if isinstance(label, str) and re.fullmatch(r'重仓[1-5]', label):
            definition = ('holdingStock' + label[-1], label, 'text', '')
        if not definition:
            continue
        key, title, kind, unit = definition
        cell = sheet.cell(row, column)
        formula = formulas[sheet.title].cell(row, column)
        value = cell.value
        status = 'ok'
        if cell.data_type == 'e' or (isinstance(value, str) and (value.startswith(('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')) or any(word in value for word in ['抓取中', '抓取失败', '加载失败', '请求失败']))):
            status = 'error'
        elif value is None or str(value).strip() in ['', '-', '—']:
            status = 'error' if formula.data_type == 'f' and '#REF!' in formula.value else 'missing'
        elif isinstance(value, dt.datetime):
            value = value.date().isoformat()
        elif kind == 'number' and not isinstance(value, (int, float)):
            try:
                value = float(str(value).strip())
            except ValueError:
                status = 'error'
        report_date = label.date().isoformat() if isinstance(label, dt.datetime) else None
        if not report_date and formula.data_type == 'f':
            date_match = re.search(r'"(\d{4}-\d{2}-\d{2})"', formula.value)
            report_date = date_match.group(1) if date_match else None
        field_catalog[key] = {'key': key, 'label': title, 'type': kind, 'unit': unit}
        result[key] = {'value': value if status == 'ok' else None, 'status': status, 'raw': str(cell.value) if status == 'error' else None,
                       'cell': cell.coordinate, 'reportDate': report_date}
    return result
evidence = collections.defaultdict(list)
scanned = []
names = {}
for sheet in book:
    section = None
    count = 0
    for row in sheet:
        name = str(row[0].value or '').strip()
        code = code_of(row[1].value)
        if not code:
            if name and (row[1].value == '基金代码' or name in ['高波固收+', '中波固收+', '低波固收+', '极低波固收+', '固收+']):
                section = name
            continue
        names.setdefault(code, name)
        count += 1
        explicit = [str(cell.value).strip() for cell in row[23:26] if str(cell.value).strip() in ['极低波', '低波', '中波', '高波']]
        attribute = explicit[0] if len(set(explicit)) == 1 else None
        if not attribute and section and '纯债' in section:
            attribute = '原表纯债分组'
        evidence[code].append({'sheet': sheet.title, 'row': row[0].row, 'section': section, 'attribute': attribute, 'name': name})
    scanned.append({'sheet': sheet.title, 'fundRows': count, 'role': '主名单（用户确认榜单）' if sheet.title == '求赚固收+' else '主名单' if sheet.title in ['求赚安稳盈', '求稳货架', '风格化固收+'] else '补充核验，不整表累加'})

def add(sheet, row, shelf, sub, state, raw_shelf, raw_class, issue=None):
    code = code_of(sheet.cell(row, 2).value)
    if not code:
        return
    name = str(sheet.cell(row, 1).value).strip()
    market_cell = {'求赚安稳盈': 'X1', '求稳货架': 'Z2', '风格化固收+': 'AA1'}.get(sheet.title)
    cached_date = sheet[market_cell].value if market_cell else None
    relations.append({
        'code': code, 'name': name, 'company': company(name), 'shelf': shelf, 'sub': sub,
        'status': state if state in ['留池', '入池', '出池', '名单收录'] else '待核实',
        'source': {'sheet': sheet.title, 'row': row, 'rawShelf': raw_shelf, 'rawClass': raw_class, 'rawStatus': state, 'dataDate': None,
                   'rawSection': next((item['section'] for item in evidence[code] if item['sheet'] == sheet.title and item['row'] == row), None),
                   'marketDate': cached_date.date().isoformat() if isinstance(cached_date, dt.datetime) else None,
                   'marketDateFormula': formulas[sheet.title][market_cell].value if market_cell else None},
        'issue': issue,
        'fields': extract_fields(sheet, row),
    })

sheet = book['求赚安稳盈']
for row in range(2, 165):
    if not code_of(sheet.cell(row, 2).value):
        continue
    raw_shelf = sheet.cell(row, 31).value
    sub = raw_shelf if raw_shelf in ['安理财', '稳理财', '盈理财'] else None
    add(sheet, row, '稳健理财' if sub else None, sub, sheet.cell(row, 30).value, raw_shelf, None,
        None if sub else '原表未标注所属货架')

sheet = book['求稳货架']
for row in range(3, 35):
    if not code_of(sheet.cell(row, 2).value):
        continue
    raw_shelf = sheet.cell(row, 33).value
    raw_class = str(sheet.cell(row, 4).value or '').strip()
    match = re.search(r'金选收益\+（(.*?)）', str(raw_shelf))
    listed_class = match.group(1) if match else None
    conflict = listed_class and raw_class and listed_class != raw_class
    sub = listed_class or (raw_class if raw_class in ['年年', '季季', '月月'] else None)
    add(sheet, row, '金选收益+', sub, sheet.cell(row, 32).value, raw_shelf, raw_class)
    if conflict:
        relations[-1]['multipleSubclassesConfirmed'] = True
        relations[-1]['source']['classificationBasis'] = '用户确认两个分类字段代表同时覆盖两个二级榜单'
        add(sheet, row, '金选收益+', raw_class, sheet.cell(row, 32).value, raw_shelf, raw_class)
        relations[-1]['multipleSubclassesConfirmed'] = True
        relations[-1]['source']['classificationBasis'] = '用户确认两个分类字段代表同时覆盖两个二级榜单'
for row in range(45, 54):
    add(sheet, row, '金选理财+', None, sheet.cell(row, 31).value, sheet.cell(row, 32).value, None)

sheet = book['风格化固收+']
section = None
for row in range(1, 24):
    if sheet.cell(row, 1).value in ['科技', '红利', '量化', '资源']:
        section = sheet.cell(row, 1).value
    if code_of(sheet.cell(row, 2).value):
        add(sheet, row, '风格化固收+', section, '名单收录', '风格化固收+', section)

sheet = book['求赚固收+']
section = None
for row in range(1, sheet.max_row + 1):
    title = sheet.cell(row, 1).value
    if title in ['极低波固收+', '低波固收+', '中波固收+', '高波固收+']:
        section = title
    if section and code_of(sheet.cell(row, 2).value):
        add(sheet, row, '金选固收+', section.replace('固收+', ''), '名单收录', None, section)
        relations[-1]['source']['rawStatus'] = None
        relations[-1]['source']['inclusionBasis'] = '用户确认本工作表四个分组作为金选固收+榜单，含极低波；非原表入池状态'
        relations[-1]['multipleSubclassesConfirmed'] = True
        relations[-1]['source']['classificationBasis'] = '用户确认保留原表各二级分组收录，同一基金可同时出现在低波与极低波榜单'
for item in relations:
    groups = {other['sub'] for other in relations if other['code'] == item['code'] and other['shelf'] == item['shelf']}
    if len(groups) > 1 and not item.get('multipleSubclassesConfirmed'):
        item['issue'] = '同一基金代码在同一货架出现多个二级分类，归属待核实；一级席位只计一次'

products = []
for code in sorted({item['code'] for item in relations}):
    items = evidence[code]
    attributes = {item['attribute'] for item in items if item['attribute']}
    products.append({'id': code, 'code': code, 'name': names[code], 'company': company(names[code]),
        'attribute': next(iter(attributes)) if len(attributes) == 1 else '未分类',
        'attributeConflict': len(attributes) > 1, 'attributeEvidence': [item for item in items if item['attribute']],
        'nameVariants': list(dict.fromkeys(item['name'] for item in items)),
        'companyBasis': '原表产品名称前缀与别名归一；管理人主体未联网核验',
        'mainProductId': None, 'shareMappingStatus': '暂无可靠主产品映射，按基金代码去重'})

confirmed = [item for item in relations if item['status'] in ['入池', '留池', '名单收录'] and item['shelf']]
counts = collections.Counter(item['shelf'] for item in confirmed)
checks = [{'label': key, 'actual': value} for key, value in counts.items()]
primary = [item for item in confirmed if item['source']['sheet'] in ['求赚安稳盈', '求稳货架']]
style = [item for item in confirmed if item['shelf'] == '风格化固收+']
audit = {
    'primaryCodes': len({item['code'] for item in primary}),
    'combinedCodes': len({item['code'] for item in confirmed}),
    'combinedSeats': len({(item['code'], item['shelf']) for item in confirmed}),
    'styleCodes': len({item['code'] for item in style}),
    'styleCompanies': len({item['company'] for item in style if item['company']}),
    'goldCodes': len({item['code'] for item in confirmed if item['shelf'] == '金选固收+'}),
    'primarySubcounts': dict(collections.Counter(item['sub'] for item in primary if item['shelf'] == '稳健理财')),
}
rules = []
rule_sheet = book['求稳货架']
for rule_row in [61, 62, 63, 69, 70, 71]:
    header_row = 60 if rule_row < 67 else 68
    cells = []
    for rule_column in range(1, 9):
        title = rule_sheet.cell(header_row, rule_column).value
        value = rule_sheet.cell(rule_row, rule_column).value
        if title is None or value is None or str(value).strip() in ['—', '-', '']:
            continue
        if isinstance(value, (int, float)) and '%' in rule_sheet.cell(rule_row, rule_column).number_format:
            value = f'{value * 100:g}%'
        cells.append({'label': str(title), 'value': str(value), 'cell': rule_sheet.cell(rule_row, rule_column).coordinate})
    rules.append({'shelf': '金选收益+' if rule_row < 67 else '金选理财+', 'sub': None,
                  'appliesTo': rule_sheet.cell(rule_row, 1).value, 'fields': cells,
                  'source': {'sheet': rule_sheet.title, 'row': rule_row, 'date': None},
                  'effectiveDate': None, 'suggestedHolding': None,
                  'uncertainty': '原表“>=85”缺少百分号，按原文展示，待核实' if rule_row == 62 else None})

payload = {'meta': {'sourceFile': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'importedAt': dt.datetime.now(dt.timezone.utc).isoformat(), 'listDate': None,
    'scope': '工作簿收录快照（跨来源名单日期未确认）', 'identityBasis': '基金代码去重，未合并份额',
    'dateNote': 'TODAY等公式日期是行情取数日，不是名单生效日。合并视图仅为工作簿阶段性汇总，不代表同一时点的在架名单。',
    'goldIncomeBasis': '求稳货架收益+区块明确列示，状态列确认；用户确认两个分类字段不一致时按同时覆盖两个二级榜单记录。一级产品与席位仍去重。',
    'goldFixedIncomeConfirmed': True,
    'goldFixedIncomeStatus': '用户确认：求赚固收+的极低波、低波、中波、高波四组作为金选固收+名单；同一基金可覆盖多个二级榜单，各二级组内去重，一级产品和席位去重。',
    'scannedSheets': scanned}, 'products': products, 'relations': relations, 'audit': audit, 'fieldCatalog': field_catalog, 'rules': rules}
target = root / 'public' / 'data' / 'shelves.json'
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
target.with_suffix('.json.gz').write_bytes(gzip.compress(target.read_bytes(), mtime=0))
print(json.dumps({'audit': audit, 'unmatchedCompanies': [item['name'] for item in products if not item['company']], 'relations': len(relations)}, ensure_ascii=False))
